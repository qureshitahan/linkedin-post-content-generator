"""Retrieve relevant snippets from a principle's indexed documents."""

import io
import logging
import re
from pathlib import Path
from typing import List, Optional, Tuple

from sqlalchemy.orm import Session

from app.models import Principle, PrincipleDocument
from app.services.claude import claude_service

logger = logging.getLogger(__name__)

PRINCIPLE_FILES_DIR = Path(__file__).resolve().parent.parent.parent / "principle_files"
CHUNK_SIZE = 500
TOP_K_DEFAULT = 8

SUPPORTED_EXTENSIONS = {
    "txt", "md", "csv", "pdf", "docx",
    "png", "jpg", "jpeg", "webp", "gif", "bmp", "tif", "tiff", "heic", "heif",
}

IMAGE_MEDIA_TYPES = {
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "webp": "image/webp",
    "gif": "image/gif",
    "bmp": "image/bmp",
    "tif": "image/tiff",
    "tiff": "image/tiff",
    "heic": "image/heic",
    "heif": "image/heif",
}

STOPWORDS = frozenset(
    {
        "about", "with", "from", "that", "this", "have", "been", "were", "their",
        "there", "what", "when", "where", "which", "would", "could", "should",
        "into", "your", "they", "them", "then", "than", "also", "just", "like",
        "want", "talk", "posts", "linkedin", "write", "writing",
    }
)


def _file_extension(filename: str) -> str:
    return filename.rsplit(".", 1)[-1].lower() if "." in filename else ""


def _extract_docx(content: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(content))
    parts: List[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                parts.append(" | ".join(row_cells))
    return "\n".join(parts).strip()


def _extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages).strip()


def _extract_image(content: bytes, filename: str) -> str:
    ext = _file_extension(filename)
    media_type = IMAGE_MEDIA_TYPES.get(ext)
    payload = content

    if ext not in ("jpg", "jpeg", "png", "webp", "gif"):
        try:
            from PIL import Image

            img = Image.open(io.BytesIO(content))
            if img.mode in ("RGBA", "P", "LA"):
                img = img.convert("RGB")
            elif img.mode != "RGB":
                img = img.convert("RGB")
            buf = io.BytesIO()
            img.save(buf, format="JPEG", quality=90)
            payload = buf.getvalue()
            media_type = "image/jpeg"
        except Exception as e:
            logger.warning("Image conversion failed for %s: %s", filename, e)
            if not media_type:
                raise ValueError(
                    f"Could not read image .{ext}. Try .png, .jpg, or .webp."
                ) from e

    if not media_type:
        raise ValueError(f"Unsupported image type: .{ext}")

    return claude_service.extract_text_from_image(payload, media_type, filename)


def extract_text_from_bytes(filename: str, content: bytes) -> str:
    ext = _file_extension(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(f".{e}" for e in sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"Unsupported file type: .{ext or 'unknown'}. Supported: {supported}")

    if ext in ("txt", "md", "csv"):
        return content.decode("utf-8", errors="replace").strip()

    if ext == "pdf":
        try:
            text = _extract_pdf(content)
            if text:
                return text
        except Exception as e:
            logger.warning("PDF extraction failed for %s: %s", filename, e)
        raise ValueError("Could not extract text from PDF.")

    if ext == "docx":
        try:
            text = _extract_docx(content)
            if text:
                return text
        except Exception as e:
            logger.warning("DOCX extraction failed for %s: %s", filename, e)
        raise ValueError("Could not extract text from Word document (.docx).")

    if ext in IMAGE_MEDIA_TYPES:
        try:
            text = _extract_image(content, filename)
            if text:
                return text
        except RuntimeError:
            raise
        except Exception as e:
            logger.warning("Image extraction failed for %s: %s", filename, e)
        raise ValueError("Could not extract text from image.")

    raise ValueError(f"Unsupported file type: .{ext}")


def chunk_text(text: str, chunk_size: int = CHUNK_SIZE) -> List[str]:
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    if not paragraphs:
        paragraphs = [text.strip()] if text.strip() else []

    chunks: List[str] = []
    current = ""
    for para in paragraphs:
        if len(para) > chunk_size:
            if current:
                chunks.append(current.strip())
                current = ""
            for i in range(0, len(para), chunk_size):
                piece = para[i : i + chunk_size].strip()
                if piece:
                    chunks.append(piece)
            continue
        if len(current) + len(para) + 2 <= chunk_size:
            current = f"{current}\n\n{para}".strip() if current else para
        else:
            if current:
                chunks.append(current.strip())
            current = para
    if current:
        chunks.append(current.strip())
    return chunks


def _tokenize(text: str) -> set[str]:
    return {
        w
        for w in re.findall(r"[a-z0-9][a-z0-9+\-/]{2,}", text.lower())
        if w not in STOPWORDS
    }


def _score_chunk(chunk: str, query: str) -> float:
    chunk_lower = chunk.lower()
    query_lower = query.lower()
    score = 0.0

    for phrase in re.findall(r"[a-z][a-z0-9+\-/ ]{3,40}[a-z0-9]", query_lower):
        phrase = phrase.strip()
        if len(phrase) > 4 and phrase in chunk_lower:
            score += 3.0

    query_tokens = _tokenize(query)
    chunk_tokens = _tokenize(chunk)
    overlap = query_tokens & chunk_tokens
    score += len(overlap) * 1.2

    if len(chunk) > 80:
        score += 0.3

    return score


def retrieve_snippets(
    documents: List[Tuple[str, str]],
    query: str,
    top_k: int = TOP_K_DEFAULT,
) -> List[str]:
    """Return top relevant text snippets from (filename, content) pairs."""
    if not documents or not query.strip():
        return []

    scored: List[Tuple[float, str, str]] = []
    for filename, content in documents:
        for chunk in chunk_text(content):
            s = _score_chunk(chunk, query)
            if s > 0:
                scored.append((s, filename, chunk))

    if not scored:
        for filename, content in documents:
            for chunk in chunk_text(content)[:3]:
                scored.append((0.1, filename, chunk))

    scored.sort(key=lambda x: x[0], reverse=True)
    seen: set[str] = set()
    results: List[str] = []
    for _, filename, chunk in scored:
        key = chunk[:120]
        if key in seen:
            continue
        seen.add(key)
        results.append(f"[{filename}] {chunk}")
        if len(results) >= top_k:
            break
    return results


class PrincipleRetrievalService:
    def load_documents(self, db: Session, principle_id: int) -> List[Tuple[str, str]]:
        docs = (
            db.query(PrincipleDocument)
            .filter(PrincipleDocument.principle_id == principle_id)
            .order_by(PrincipleDocument.created_at.desc())
            .all()
        )
        return [(d.filename, d.content_text) for d in docs if d.content_text.strip()]

    def retrieve(
        self,
        db: Session,
        principle_id: int,
        query: str,
        top_k: int = TOP_K_DEFAULT,
    ) -> List[str]:
        documents = self.load_documents(db, principle_id)
        return retrieve_snippets(documents, query, top_k=top_k)

    def principle_summary(self, db: Session, principle_id: int, limit: int = 15) -> str:
        """Broad context from all documents for parsing author background."""
        docs = self.load_documents(db, principle_id)
        if not docs:
            return ""
        combined_query = "experience skills achievements projects work background expertise"
        snippets = retrieve_snippets(docs, combined_query, top_k=limit)
        if not snippets:
            snippets = [f"[{name}] {text[:400]}" for name, text in docs[:5]]
        return "\n".join(snippets)

    def get_principle(self, db: Session, principle_id: int) -> Optional[Principle]:
        return db.query(Principle).filter(Principle.id == principle_id).first()

    def save_uploaded_file(self, principle_id: int, filename: str, content: bytes) -> Path:
        safe_name = re.sub(r"[^\w.\-]", "_", Path(filename).name)[:200]
        dest_dir = PRINCIPLE_FILES_DIR / str(principle_id)
        dest_dir.mkdir(parents=True, exist_ok=True)
        dest = dest_dir / safe_name
        dest.write_bytes(content)
        return dest

    def delete_file(self, stored_path: Optional[str]) -> None:
        if not stored_path:
            return
        try:
            path = Path(stored_path)
            if path.exists():
                path.unlink()
        except OSError as e:
            logger.warning("Could not delete principle file %s: %s", stored_path, e)


principle_retrieval_service = PrincipleRetrievalService()
